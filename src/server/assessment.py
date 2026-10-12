"""Private, local CV completeness and job keyword assessment."""
import base64
import io
import json
import re
import sys
import zipfile

SKILLS = {
    "Python": ["python"], "JavaScript": ["javascript", "typescript", "node.js", "nodejs"],
    "SQL": ["sql", "postgresql", "mysql", "sqlite"], "Cloud": ["aws", "azure", "gcp", "cloud"],
    "Docker": ["docker", "containers"], "Kubernetes": ["kubernetes", "k8s"],
    "React": ["react", "reactjs"], "Java": ["java"], "C# / .NET": ["c#", ".net", "asp.net"],
    "Data analysis": ["data analysis", "data analytics", "power bi", "tableau", "excel"],
    "Machine learning": ["machine learning", "tensorflow", "pytorch"], "Git": ["git", "github", "gitlab"],
    "Agile": ["agile", "scrum", "kanban"], "Project management": ["project management", "project manager", "programme management"],
    "Delivery": ["delivery manager", "delivery lead", "service delivery"],
    "Stakeholder management": ["stakeholder", "stakeholders"], "Leadership": ["leadership", "team lead", "managed a team"],
    "Communication": ["communication", "communicate", "presentation"],
    "Customer service": ["customer service", "customer support"], "Sales": ["sales", "business development"],
    "Finance": ["finance", "accounting", "budget", "financial"], "Testing": ["testing", "quality assurance", "selenium"],
    "Security": ["cybersecurity", "information security", "security"], "Design": ["ux", "ui design", "ui designer", "user interface", "figma", "design"],
    "Operations": ["operations", "process improvement", "continuous improvement"],
}


def mentions(text, term):
    return bool(re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text, re.I))


def assess(cv, jobs, minimum_words=30):
    cv = cv.strip()
    checks = [
        ("Contact email", bool(re.search(r"[^\s@]+@[^\s@]+\.[^\s@]+", cv))),
        ("Experience section", bool(re.search(r"experience|employment|work history", cv, re.I))),
        ("Skills section", bool(re.search(r"skills|competencies|expertise", cv, re.I))),
        ("Education or qualifications", bool(re.search(r"education|qualifications|certification|degree", cv, re.I))),
        ("Measurable achievements", bool(re.search(r"\d+\s*(?:%|percent)|\$\s*\d|\b(?:increased|reduced|improved|saved|delivered)\b.{0,80}\d", cv, re.I))),
    ]
    available = len(cv.split()) >= minimum_words
    overview = {"available": available, "score": sum(20 for _, found in checks if found) if available else None,
                "label": "CV completeness", "checks": [{"name": name, "present": found} for name, found in checks],
                "reason": "Add at least 30 words of readable CV text to assess your CV." if not available else "Five document checks, 20 points each; this is not a measure of employability."}
    overview["ats"] = {"readable": available,
                       "notes": ["Use selectable text, clear headings and a simple one-column document.",
                                 "Keep name and contact details in the document body, not only headers or text boxes.",
                                 "Review extracted text for missing or reordered content; visual layout is not checked here.",
                                 "ATS products parse and search documents differently. No universal ATS pass score is calculated."]}
    fits = []
    for job in jobs:
        description = job.get("description", "")
        text = job.get("role", "") + " " + description
        requested = [name for name, terms in SKILLS.items() if any(mentions(text, term) for term in terms)]
        matched = [name for name in requested if any(mentions(cv, term) for term in SKILLS[name])]
        missing = [name for name in requested if name not in matched]
        sufficient = available and len(description.split()) >= 20 and len(requested) >= 2
        evidence = []
        for name in matched:
            for line in re.split(r"\n|(?<=[.!?])\s+", cv):
                if any(mentions(line, term) for term in SKILLS[name]):
                    evidence.append({"skill": name, "text": line.strip()[:240]})
                    break
        fits.append({"score": round(100 * len(matched) / len(requested)) if sufficient else None,
                     "label": "Keyword match", "matched": matched, "missing": missing,
                     "evidence": evidence, "limited": not sufficient,
                     "reason": "Save readable CV text first." if not available else
                     "The listing needs a fuller description or more identifiable skills." if not sufficient else
                     f"Your CV mentions {len(matched)} of {len(requested)} recognized skill topics in the listing. Review requirements and evidence; no experience level or eligibility is verified."})
    skills = [name for name, terms in SKILLS.items() if any(mentions(cv, term) for term in terms)]
    roles = {
        "Project management": "Project Manager", "Delivery": "Delivery Lead",
        "Python": "Python Developer", "JavaScript": "Software Engineer",
        "SQL": "Data Analyst", "Leadership": "Team Lead",
        'Cloud': 'Cloud Engineer', 'Docker': 'DevOps Engineer', 'Kubernetes': 'DevOps Engineer',
        'React': 'Frontend Developer', 'Java': 'Java Developer', 'C# / .NET': '.NET Developer',
        'Data analysis': 'Data Analyst', 'Machine learning': 'Machine Learning Engineer',
        'Agile': 'Scrum Master', 'Stakeholder management': 'Business Analyst',
        'Customer service': 'Customer Service Representative', 'Sales': 'Sales Representative',
        'Finance': 'Financial Analyst', 'Testing': 'QA Engineer', 'Security': 'Security Analyst',
        'Design': 'UX Designer', 'Operations': 'Operations Manager',
    }
    profession_titles = ['Registered Nurse', 'Nurse', 'Teacher', 'Lecturer', 'Civil Engineer',
                         'Mechanical Engineer', 'Electrical Engineer', 'Accountant', 'Architect',
                         'Lawyer', 'Doctor', 'Pharmacist', 'Physiotherapist', 'Chef', 'Electrician',
                         'Plumber', 'Marketing Manager', 'HR Manager', 'Administrative Assistant',
                         'Research Scientist', 'Warehouse Manager', 'Truck Driver']
    inferred = [title for title in profession_titles if mentions(cv, title)]
    recommendations = list(dict.fromkeys(inferred + [roles[skill] for skill in skills if skill in roles]))[:6] if available else []
    return {"cv": overview, "fits": fits, "recommendations": recommendations, "skills": skills}


def extract(kind, encoded):
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
    except ImportError:
        pass
    contents = base64.b64decode(encoded, validate=True)
    if kind == "txt":
        text = contents.decode("utf-8-sig", errors="replace")
    elif kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(contents))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("Use an unencrypted PDF with up to 100 pages.")
        text = "\n".join((page.extract_text() or "")[:100000] for page in reader.pages)
    elif kind == "docx":
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Word document is too large to extract.")
        document = Document(io.BytesIO(contents))
        text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    else:
        raise ValueError("Unsupported document format.")
    return text[:100000].strip()


if __name__ == "__main__":
    try:
        data = json.load(sys.stdin)
        if data.get('interview'):
            from interview import prepare
            result = prepare(data.get('cv', ''), data.get('job', {}))
        elif data.get('recommendDocuments'):
            from document_assessment import recommend_documents
            result = recommend_documents(data.get('cv', ''), data.get('cover', ''))
        elif data.get('coverAssessment'):
            from document_assessment import cover_assessment
            result = cover_assessment(data.get('text', ''), data.get('job'))
        elif data.get('extract') and data.get('purpose'):
            contents = extract(data['kind'], data['base64'])
            from document_assessment import cover_assessment
            result = {'text': contents, 'assessment': cover_assessment(contents, data.get('job')) if data['purpose'] == 'Cover letter' else assess(contents, [])['cv']}
        else:
            result = {"text": extract(data["kind"], data["base64"])} if data.get("extract") else assess(data.get("cv", ""), data.get("jobs", []))
        print(json.dumps(result))
    except Exception:
        print("CV processing failed. Use reviewed text or a readable PDF/DOCX/TXT.", file=sys.stderr)
        sys.exit(1)
