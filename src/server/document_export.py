"""Template-aware DOCX export and isolated, local LibreOffice PDF conversion."""
import base64
from copy import deepcopy
import io
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

def create_document(template):
    from docx import Document
    from docx.shared import Pt, RGBColor
    raw = base64.b64decode(template.get('base64', ''), validate=True)
    kind = template.get('kind')
    if kind == 'docx':
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('Template is too large.')
        return Document(io.BytesIO(raw)), []
    document = Document()
    document.styles['Normal'].font.name = 'Liberation Sans'
    document.styles['Normal'].font.size = Pt(11)
    warnings = ['No editable template: a readable document layout is used.']
    if kind == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted or not reader.pages or len(reader.pages) > 100:
            raise ValueError('Use an unencrypted, readable PDF template.')
        page = reader.pages[0]
        section = document.sections[0]
        section.page_width = Pt(float(page.mediabox.width))
        section.page_height = Pt(float(page.mediabox.height))
        spans = []
        current_color = [0, 0, 0]
        def operand(operator, operands, cm, tm):
            if operator == b'rg':
                current_color[:] = [float(x) for x in operands[:3]]
            elif operator == b'g':
                current_color[:] = [float(operands[0])] * 3
        def visitor(text, cm, tm, font, size):
            if text.strip() and font:
                spans.append((str(font.get('/BaseFont', '/Helvetica')).split('+')[-1].lstrip('/'), float(size), len(text), tuple(current_color), float(tm[4]), float(tm[5])))
        page.extract_text(visitor_text=visitor, visitor_operand_before=operand)
        if spans:
            font, size, _, color, _, _ = max(spans, key=lambda s: s[2])
            document.styles['Normal'].font.name = font.replace('-Bold', '').replace('-Italic', '')
            document.styles['Normal'].font.size = Pt(max(8, min(24, size)))
            document.styles['Normal'].font.color.rgb = RGBColor(*[round(max(0, min(1, x)) * 255) for x in color])
            largest = max(spans, key=lambda s: s[1])
            for name in ('Title', 'Heading 1'):
                document.styles[name].font.size = Pt(max(12, min(32, largest[1])))
                document.styles[name].font.color.rgb = RGBColor(*[round(max(0, min(1, x)) * 255) for x in largest[3]])
            xs = [s[4] for s in spans if 5 < s[4] < float(page.mediabox.width)/2]
            if xs:
                section.left_margin = section.right_margin = Pt(max(24, min(108, min(xs))))
        warnings = ['PDF page size and detected typography are matched. Fixed PDF columns, graphics and exact positioning cannot be reliably reconstructed; upload the original DOCX for template preservation.']
    return document, warnings

def replace_body(document, text):
    from difflib import SequenceMatcher
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    lines = [re.sub(r'^#{1,6}\s+', '', line).strip() for line in text.splitlines() if line.strip()]
    slots = [Paragraph(e, document) for e in document.element.body.iter(qn('w:p')) if Paragraph(e, document).text.strip()]
    if not slots:
        for index, line in enumerate(lines):
            heading = len(line) < 60 and (line.isupper() or line.lower().rstrip(':') in ('experience', 'skills', 'education', 'summary', 'qualifications', 'certifications'))
            document.add_paragraph(line, style='Title' if index == 0 and len(line) < 60 else 'Heading 1' if heading else 'Normal')
        return

    def key(value):
        value = re.sub(r'[^a-z ]', '', value.casefold()).strip()
        return {'professional experience': 'experience', 'work experience': 'experience', 'professional summary': 'summary', 'profile': 'summary', 'technical skills': 'skills', 'core skills': 'skills'}.get(value, value)

    def write(paragraph, value):
        # Map matching characters back to their original runs, retaining mixed emphasis.
        runs = list(paragraph.runs)
        original = paragraph.text
        formats = []
        for run in runs:
            formats.extend([run._r.rPr] * len(run.text))
        if not formats:
            formats = [None] * max(1, len(original))
        style_at = [None] * len(value)
        matcher = SequenceMatcher(None, original, value, autojunk=False)
        for tag, a, b, c, d in matcher.get_opcodes():
            for i in range(c, d):
                source = a + (i - c) if tag == 'equal' else min(a, len(formats) - 1)
                style_at[i] = formats[min(source, len(formats) - 1)]
        for child in list(paragraph._p):
            if child.tag == qn('w:pPr'):
                continue
            if list(child.iter(qn('w:drawing'))) or list(child.iter(qn('w:pict'))):
                for node in child.iter(qn('w:t')):
                    node.text = ''
            else:
                paragraph._p.remove(child)
        offset = 0
        while offset < len(value):
            end = offset + 1
            while end < len(value) and style_at[end] is style_at[offset]:
                end += 1
            run = paragraph.add_run(value[offset:end])
            if style_at[offset] is not None:
                run._r.insert(0, deepcopy(style_at[offset]))
            offset = end

    def fill(paragraphs, values, section=True):
        if not paragraphs:
            return
        # Insert overflow next to its section, inside the same table cell/column.
        prototype = next((p for p in reversed(paragraphs) if not p.style.name.lower().startswith(('heading', 'title'))), paragraphs[-1])
        anchor = paragraphs[-1]._p
        if not section:
            prototype = next((p for p in reversed(paragraphs) if p._p.getparent() is document.element.body and not p.style.name.lower().startswith(('heading', 'title'))), paragraphs[0])
            anchor = document.element.body[-2] if document.element.body[-1].tag == qn('w:sectPr') else document.element.body[-1]
        for index, value in enumerate(values):
            if index < len(paragraphs):
                paragraph = paragraphs[index]
            else:
                cloned = deepcopy(prototype._p)
                # A cloned text paragraph must not duplicate pictures or bookmarks.
                for tag in ('w:drawing', 'w:pict', 'w:bookmarkStart', 'w:bookmarkEnd'):
                    for node in list(cloned.iter(qn(tag))):
                        node.getparent().remove(node)
                anchor.addnext(cloned)
                anchor = cloned
                paragraph = Paragraph(cloned, prototype._parent)
            write(paragraph, value)
        for paragraph in paragraphs[len(values):]:
            write(paragraph, '')

    anchors = [(i, key(p.text)) for i, p in enumerate(slots) if 'heading' in p.style.name.lower() or key(p.text) in ('experience', 'skills', 'education', 'summary', 'qualifications', 'certifications', 'projects')]
    known = {name for _, name in anchors}
    draft_anchors = [(i, key(line)) for i, line in enumerate(lines) if key(line) in known]
    if not anchors or not draft_anchors:
        fill(slots, lines, section=False)
        return
    sections = {}
    for position, (index, name) in enumerate(draft_anchors):
        stop = draft_anchors[position + 1][0] if position + 1 < len(draft_anchors) else len(lines)
        sections.setdefault(name, []).extend(lines[index + 1:stop])
    preamble = slots[:anchors[0][0]]
    preamble_lines = lines[:draft_anchors[0][0]]
    if not preamble and preamble_lines:
        cloned = deepcopy(slots[0]._p)
        slots[0]._p.addprevious(cloned)
        paragraph = Paragraph(cloned, slots[0]._parent)
        paragraph.style = document.styles['Normal']
        preamble = [paragraph]
    fill(preamble, preamble_lines)
    for position, (index, name) in enumerate(anchors):
        stop = anchors[position + 1][0] if position + 1 < len(anchors) else len(slots)
        values = sections.pop(name, [])
        paragraphs = slots[index + 1:stop]
        if not paragraphs:
            # Heading-only sections gain a body paragraph immediately below the heading.
            cloned = deepcopy(slots[index]._p)
            slots[index]._p.addnext(cloned)
            paragraph = Paragraph(cloned, slots[index]._parent)
            paragraph.style = document.styles['Normal']
            paragraphs = [paragraph]
        fill(paragraphs, values)
    # Repeated headings are kept in their original locations; no draft text is lost.

def render(data):
    text = data.get('text', '').strip()
    if not text or len(text) > 100000 or data.get('format') not in ('docx', 'pdf'):
        raise ValueError('Enter a draft and choose DOCX or PDF.')
    document, warnings = create_document(data.get('template') or {})
    replace_body(document, text)
    # Disable external template resources before office conversion.
    for part in list(document.part.package.parts):
        for key, relationship in list(part.rels.items()):
            if relationship.is_external:
                del part.rels[key]
    from docx.oxml.ns import qn
    for element in list(document.element.iter(qn('w:hyperlink'))):
        parent = element.getparent()
        position = parent.index(element)
        for child in list(element):
            parent.insert(position, child)
            position += 1
        parent.remove(element)
    output = io.BytesIO()
    document.save(output)
    if data['format'] == 'docx':
        return {'base64': base64.b64encode(output.getvalue()).decode(), 'warnings': warnings}
    with tempfile.TemporaryDirectory(prefix='milo-export-') as directory:
        root = Path(directory)
        source = root / 'draft.docx'
        source.write_bytes(output.getvalue())
        profile = (root / 'office-profile').as_uri()
        profile_dir = root / 'office-profile' / 'user'
        profile_dir.mkdir(parents=True)
        (profile_dir / 'registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding='utf-8')
        subprocess.run(['soffice', '-env:UserInstallation=' + profile, '--headless', '--norestore', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(root), str(source)], check=True, timeout=45, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = root / 'draft.pdf'
        if not pdf.exists() or pdf.stat().st_size > 20 * 1024 * 1024:
            raise ValueError('PDF conversion failed.')
        return {'base64': base64.b64encode(pdf.read_bytes()).decode(), 'warnings': warnings + ['Review PDF pagination and fonts; unavailable proprietary fonts can be substituted by the local renderer.']}

if __name__ == '__main__':
    try:
        print(json.dumps(render(json.load(sys.stdin))))
    except Exception:
        print('Document export failed. Check the template and try again.', file=sys.stderr)
        sys.exit(1)
