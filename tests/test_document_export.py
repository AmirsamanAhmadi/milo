import base64
import importlib.util
import io
from pathlib import Path
import shutil
import sys
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src/server'))
from document_export import render

@unittest.skipUnless(importlib.util.find_spec('docx'), 'Export dependencies are checked in the runtime image')
class ExportTests(unittest.TestCase):
    def template(self):
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        d=Document()
        d.sections[0].left_margin=Inches(.7)
        d.sections[0].header.paragraphs[0].text='Original header'
        d.styles['Normal'].font.name='Liberation Sans'
        d.styles['Normal'].font.size=Pt(12)
        p=d.add_paragraph('Original title',style='Title')
        p.runs[0].font.color.rgb=RGBColor(20,70,140)
        d.add_paragraph('Original body')
        table=d.add_table(rows=1,cols=1)
        table.cell(0,0).text='Original table text'
        output=io.BytesIO();d.save(output)
        return {'kind':'docx','base64':base64.b64encode(output.getvalue()).decode()}

    def test_docx_preserves_template_styles_headers_tables_and_all_draft_text(self):
        from docx import Document
        text='New title\nNew evidence\nNew table content\nAdditional achievement'
        result=render({'text':text,'format':'docx','template':self.template()})
        document=Document(io.BytesIO(base64.b64decode(result['base64'])))
        self.assertEqual(document.sections[0].left_margin.inches,.7)
        self.assertEqual(document.sections[0].header.paragraphs[0].text,'Original header')
        self.assertEqual(document.paragraphs[0].style.name,'Title')
        self.assertEqual(str(document.paragraphs[0].runs[0].font.color.rgb),'14468C')
        self.assertEqual(document.tables[0].cell(0,0).text,'New table content')
        self.assertIn('Additional achievement',[p.text for p in document.paragraphs])
        self.assertNotIn('Original body',[p.text for p in document.paragraphs])

    @unittest.skipUnless(shutil.which('soffice'), 'PDF renderer is checked in the runtime image')
    def test_pdf_export_has_readable_draft_content(self):
        from pypdf import PdfReader
        result=render({'text':'New title\nPython delivery improved by 35%.\nSQL experience','format':'pdf','template':self.template()})
        data=base64.b64decode(result['base64'])
        self.assertTrue(data.startswith(b'%PDF-'))
        text=' '.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        self.assertIn('35%',text)
        self.assertIn('Original header',text)
        self.assertNotIn('Original body',text)

    def test_missing_template_is_disclosed_and_empty_draft_is_rejected(self):
        self.assertTrue(render({'text':'Plain draft','format':'docx'})['warnings'])
        with self.assertRaises(ValueError):render({'text':'','format':'docx'})

    def test_reordered_sections_and_extra_lines_stay_in_original_table_cells(self):
        from docx import Document
        d=Document()
        d.add_paragraph('Ada Example',style='Title')
        table=d.add_table(rows=1,cols=2)
        left,right=table.rows[0].cells
        left.paragraphs[0].text='Skills'
        left.paragraphs[0].style='Heading 1'
        left.add_paragraph('Python and SQL')
        right.paragraphs[0].text='Experience'
        right.paragraphs[0].style='Heading 1'
        p=right.add_paragraph()
        p.add_run('Engineer').bold=True
        p.add_run(' at Example Ltd').italic=True
        right.add_paragraph('Built an API')
        d.add_paragraph('Education',style='Heading 1')
        d.add_paragraph('BSc Computing')
        output=io.BytesIO();d.save(output)
        result=render({'text':'Ada Example\nExperience\nEngineer at Example Ltd\nBuilt a secure API\nReduced latency by 30%\nSkills\nPython and SQL\nDocker\nEducation\nBSc Computing','format':'docx','template':{'kind':'docx','base64':base64.b64encode(output.getvalue()).decode()}})
        exported=Document(io.BytesIO(base64.b64decode(result['base64'])))
        left,right=exported.tables[0].rows[0].cells
        self.assertEqual(left.text,'Skills\nPython and SQL\nDocker')
        self.assertEqual(right.text,'Experience\nEngineer at Example Ltd\nBuilt a secure API\nReduced latency by 30%')
        self.assertTrue(right.paragraphs[1].runs[0].bold)
        self.assertTrue(right.paragraphs[1].runs[-1].italic)
        self.assertEqual(exported.paragraphs[-2].style.name,'Heading 1')
        self.assertEqual(exported.paragraphs[-1].text,'BSc Computing')
